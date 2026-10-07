#!/usr/bin/env python3
"""Build the blank revision sheets that /revision-timetable/ offers as downloads.

    python3 scripts/make-revision-sheets.py

writes three files into static/downloads/, which Hugo copies into the build:

  revision-timetable-weekly.pdf    A4 landscape week to fill in by hand
  revision-topic-checklist.pdf     A4 portrait red, amber, green topic list
  revision-timetable.xlsx          both sheets again, editable, for Excel or
                                   Google Sheets

Run it again after changing the wording below, and commit the new files. The
PDFs come out byte-for-byte the same on every run, so an unchanged sheet never
shows up as a change in git.

Needs reportlab and openpyxl: pip3 install reportlab openpyxl
"""
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from reportlab import rl_config
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen import canvas
except ImportError:
    raise SystemExit("make-revision-sheets.py needs reportlab and openpyxl: "
                     "pip3 install reportlab openpyxl")

rl_config.invariant = 1  # no timestamps or random IDs in the PDFs

OUT = Path(__file__).resolve().parent.parent / "static" / "downloads"
PAGE_URL = "thedegreegap.com/study/revision-timetable"

BURGUNDY = "#800020"
TEXT = "#171717"
MUTED = "#5f5a54"
LINE = "#b9afa3"
CREAM = "#fbf7ef"
RAG = {"R": "#f6d5d5", "A": "#fbe7bf", "G": "#d6ecdc"}

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
WEEKLY_INTRO = ("Write a time down the side, then a subject and a task in each box. "
                "Sessions of 30 to 45 minutes with a short break between work best. "
                "Tick the corner when a session is done.")
WEEKLY_TIPS = ["Test yourself rather than re-reading.",
               "Mix two or three subjects in a day.",
               "Do past paper questions against the clock.",
               "Keep one evening a week free."]
CHECKLIST_INTRO = ("Copy the topics from your specification or your teacher's list. "
                   "Rate each one R (can't do it yet), A (shaky) or G (confident). "
                   "Start with the reds, and rate a topic again after every test or past paper.")
CHECKLIST_COLUMNS = ["Topic", "R", "A", "G", "Revised on", "Tested again on", "Past paper done"]


def wrap(c, text, font, size, width):
    """Split text into lines no wider than width at this font size."""
    lines, line = [], ""
    for word in text.split():
        trial = f"{line} {word}".strip()
        if c.stringWidth(trial, font, size) <= width:
            line = trial
        else:
            lines.append(line)
            line = word
    return lines + [line] if line else lines


def header(c, w, h, m, title, right_label):
    c.setFillColor(HexColor(BURGUNDY))
    c.setFont("Helvetica-Bold", 20)
    c.drawString(m, h - m - 16, title)
    c.setFillColor(HexColor(TEXT))
    c.setFont("Helvetica", 10.5)
    label_w = c.stringWidth(right_label, "Helvetica", 10.5)
    c.drawString(w - m - 150 - label_w - 6, h - m - 14, right_label)
    c.setStrokeColor(HexColor(LINE))
    c.setLineWidth(0.8)
    c.line(w - m - 150, h - m - 16, w - m, h - m - 16)


def footer(c, w, m):
    c.setFont("Helvetica", 8)
    c.setFillColor(HexColor(MUTED))
    c.drawString(m, m - 12, f"Make a personalised timetable for free at {PAGE_URL}")
    c.setFont("Helvetica-Bold", 8)
    c.setFillColor(HexColor(BURGUNDY))
    c.drawRightString(w - m, m - 12, "The Degree Gap")


def weekly_pdf(path):
    w, h = landscape(A4)
    m = 28
    c = canvas.Canvas(str(path), pagesize=(w, h))
    c.setTitle("Weekly revision timetable")
    c.setAuthor("The Degree Gap")
    c.setSubject("A blank weekly revision timetable to print and fill in")
    header(c, w, h, m, "My revision timetable", "Week starting")

    c.setFont("Helvetica", 9.5)
    c.setFillColor(HexColor(MUTED))
    y = h - m - 36
    for line in wrap(c, WEEKLY_INTRO, "Helvetica", 9.5, w - 2 * m):
        c.drawString(m, y, line)
        y -= 12

    # The grid: a narrow time column, then the seven days.
    top = y - 6
    time_w = 58
    day_w = (w - 2 * m - time_w) / 7
    head_h, rows, row_h = 22, 8, 41
    c.setFillColor(HexColor(BURGUNDY))
    c.rect(m, top - head_h, w - 2 * m, head_h, stroke=0, fill=1)
    c.setFillColor(HexColor("#ffffff"))
    c.setFont("Helvetica-Bold", 10)
    c.drawCentredString(m + time_w / 2, top - 15, "Time")
    for i, day in enumerate(DAYS):
        c.drawCentredString(m + time_w + day_w * (i + 0.5), top - 15, day)
    grid_top = top - head_h
    grid_bottom = grid_top - rows * row_h
    c.setFillColor(HexColor(CREAM))
    c.rect(m + time_w + day_w * 5, grid_bottom, day_w * 2, rows * row_h, stroke=0, fill=1)
    c.setStrokeColor(HexColor(LINE))
    c.setLineWidth(0.6)
    for r in range(rows + 1):
        c.line(m, grid_top - r * row_h, w - m, grid_top - r * row_h)
    c.line(m, grid_top, m, grid_bottom)
    for i in range(8):
        x = m + time_w + day_w * i
        c.line(x, grid_top, x, grid_bottom)
    # A small tick box in the corner of every session box.
    c.setLineWidth(0.5)
    for r in range(rows):
        for i in range(7):
            x = m + time_w + day_w * (i + 1) - 11
            c.rect(x, grid_top - (r + 1) * row_h + 4, 7, 7, stroke=1, fill=0)

    # Three boxes under the grid.
    box_top = grid_bottom - 12
    box_h = box_top - (m + 4)
    gap = 12
    box_w = (w - 2 * m - 2 * gap) / 3
    titles = ["This week's focus", "Subject colour key", "Remember"]
    for i, title in enumerate(titles):
        x = m + i * (box_w + gap)
        c.setStrokeColor(HexColor(LINE))
        c.setLineWidth(0.6)
        c.roundRect(x, box_top - box_h, box_w, box_h, 4, stroke=1, fill=0)
        c.setFillColor(HexColor(BURGUNDY))
        c.setFont("Helvetica-Bold", 10)
        c.drawString(x + 10, box_top - 16, title)
        c.setFillColor(HexColor(TEXT))
        if i == 0:
            for k in range(4):
                ly = box_top - 36 - k * 18
                c.line(x + 10, ly, x + box_w - 10, ly)
        elif i == 1:
            for k in range(8):
                col, row = divmod(k, 4)
                kx = x + 10 + col * (box_w / 2)
                ky = box_top - 34 - row * 18
                c.rect(kx, ky - 1, 9, 9, stroke=1, fill=0)
                c.line(kx + 15, ky - 1, kx + box_w / 2 - 22, ky - 1)
        else:
            c.setFont("Helvetica", 9.5)
            for k, tip in enumerate(WEEKLY_TIPS):
                c.drawString(x + 10, box_top - 34 - k * 15, f"•  {tip}")

    footer(c, w, m)
    c.showPage()
    c.save()


def checklist_pdf(path):
    w, h = A4
    m = 28
    c = canvas.Canvas(str(path), pagesize=(w, h))
    c.setTitle("Revision topic checklist")
    c.setAuthor("The Degree Gap")
    c.setSubject("A red, amber, green checklist for revising one subject")
    header(c, w, h, m, "Revision topic checklist", "Subject")

    c.setFont("Helvetica", 10.5)
    c.setFillColor(HexColor(TEXT))
    y = h - m - 40
    fields = [("Exam board", 150), ("Paper", 150)]
    x = m
    for label, line_w in fields:
        c.drawString(x, y, label)
        lx = x + c.stringWidth(label, "Helvetica", 10.5) + 6
        c.line(lx, y - 2, lx + line_w, y - 2)
        x = lx + line_w + 24

    c.setFont("Helvetica", 9.5)
    c.setFillColor(HexColor(MUTED))
    y -= 22
    for line in wrap(c, CHECKLIST_INTRO, "Helvetica", 9.5, w - 2 * m):
        c.drawString(m, y, line)
        y -= 12

    widths = [205, 26, 26, 26, 80, 92, 84]
    top = y - 6
    head_h, rows, row_h = 22, 28, 23.5
    c.setFillColor(HexColor(BURGUNDY))
    c.rect(m, top - head_h, sum(widths), head_h, stroke=0, fill=1)
    c.setFillColor(HexColor("#ffffff"))
    c.setFont("Helvetica-Bold", 9)
    x = m
    for col, cw in zip(CHECKLIST_COLUMNS, widths):
        if col == "Topic":
            c.drawString(x + 8, top - 15, col)
        else:
            c.drawCentredString(x + cw / 2, top - 15, col)
        x += cw
    grid_top = top - head_h
    grid_bottom = grid_top - rows * row_h
    x = m + widths[0]
    for key in "RAG":
        c.setFillColor(HexColor(RAG[key]))
        c.rect(x, grid_bottom, widths[1], rows * row_h, stroke=0, fill=1)
        x += widths[1]
    c.setStrokeColor(HexColor(LINE))
    c.setLineWidth(0.6)
    for r in range(rows + 1):
        c.line(m, grid_top - r * row_h, m + sum(widths), grid_top - r * row_h)
    x = m
    for cw in [0] + widths:
        x += cw
        c.line(x, grid_top, x, grid_bottom)

    footer(c, w, m)
    c.showPage()
    c.save()


def workbook(path):
    wb = Workbook()
    wb.properties.creator = "The Degree Gap"
    wb.properties.title = "Revision timetable and topic checklist"
    thin = Side(style="thin", color=LINE[1:])
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor=BURGUNDY[1:])
    head_font = Font(bold=True, color="FFFFFF")
    cream = PatternFill("solid", fgColor=CREAM[1:])
    title_font = Font(bold=True, size=16, color=BURGUNDY[1:])
    muted = Font(italic=True, color=MUTED[1:])
    middle = Alignment(horizontal="center", vertical="center")
    cell_wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "Weekly timetable"
    ws["A1"] = "My revision timetable"
    ws["A1"].font = title_font
    ws["F1"] = "Week starting:"
    ws["F1"].alignment = Alignment(horizontal="right")
    ws.merge_cells("G1:H1")
    ws["G1"].border = Border(bottom=thin)
    ws["H1"].border = Border(bottom=thin)
    ws["A2"] = WEEKLY_INTRO.replace("Write a time", "Type a time").replace(
        "Tick the corner when", "Colour the box in when")
    ws["A2"].font = muted
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 32
    for col, label in enumerate(["Time"] + DAYS, start=1):
        cell = ws.cell(row=4, column=col, value=label)
        cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, middle, box
    for row in range(5, 15):
        ws.row_dimensions[row].height = 42
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            cell.border, cell.alignment = box, cell_wrap
            if col >= 7:
                cell.fill = cream
    ws["A16"] = "This week's focus:"
    ws["A16"].font = Font(bold=True, color=BURGUNDY[1:])
    for row in range(17, 20):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        for col in range(1, 9):
            ws.cell(row=row, column=col).border = Border(bottom=thin)
    ws["A21"] = "Remember: " + " ".join(WEEKLY_TIPS)
    ws["A21"].font = muted
    ws["A23"] = f"Make a personalised timetable for free at {PAGE_URL}"
    ws["A23"].font = Font(size=9, color=MUTED[1:])
    ws.column_dimensions["A"].width = 12
    for letter in "BCDEFGH":
        ws.column_dimensions[letter].width = 20
    ws.freeze_panes = "B5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:H23"

    ws = wb.create_sheet("Topic checklist")
    ws["A1"] = "Revision topic checklist"
    ws["A1"].font = title_font
    for ref, label in (("A2", "Subject:"), ("C2", "Exam board:")):
        ws[ref] = label
        ws[ref].font = Font(bold=True)
    for ref in ("B2", "D2"):
        ws[ref].border = Border(bottom=thin)
    ws["A3"] = CHECKLIST_INTRO.replace("R (can't do it yet), A (shaky) or G (confident)",
                                       "Red (can't do it yet), Amber (shaky) or Green (confident)")
    ws["A3"].font = muted
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A3:E3")
    ws.row_dimensions[3].height = 46
    for col, label in enumerate(["Topic", "Rating", "Revised on", "Tested again on",
                                 "Past paper done"], start=1):
        cell = ws.cell(row=5, column=col, value=label)
        cell.fill, cell.font, cell.alignment, cell.border = head_fill, head_font, middle, box
    last = 54
    for row in range(6, last + 1):
        ws.row_dimensions[row].height = 20
        for col in range(1, 6):
            cell = ws.cell(row=row, column=col)
            cell.border = box
            if col in (3, 4):
                cell.number_format = "dd/mm/yyyy"
            if col in (2, 5):
                cell.alignment = middle
    rating = DataValidation(type="list", formula1='"Red,Amber,Green"', allow_blank=True)
    done = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(rating)
    ws.add_data_validation(done)
    rating.add(f"B6:B{last}")
    done.add(f"E6:E{last}")
    for word, colour in (("Red", RAG["R"]), ("Amber", RAG["A"]), ("Green", RAG["G"])):
        ws.conditional_formatting.add(
            f"B6:B{last}",
            CellIsRule(operator="equal", formula=[f'"{word}"'],
                       fill=PatternFill("solid", fgColor=colour[1:], bgColor=colour[1:])))
    for letter, width in zip("ABCDE", (46, 12, 14, 17, 17)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A6"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"

    wb.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    weekly_pdf(OUT / "revision-timetable-weekly.pdf")
    checklist_pdf(OUT / "revision-topic-checklist.pdf")
    workbook(OUT / "revision-timetable.xlsx")
    for name in sorted(p.name for p in OUT.iterdir()):
        print(f"  static/downloads/{name}  {(OUT / name).stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
